"""Uniform, formula-safe exports of reviewed Word extraction rows."""

import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from services.word_extraction_service import FIELDS


EXPORT_COLUMNS = {"source_file": "Source file", "source_refs": "Source locations", **FIELDS,
                  "match_status": "Database check", "matching_pids": "Matching PIDs",
                  "match_reasons": "Match reasons", "review_notes": "Review notes"}


class WordExportError(ValueError):
    pass


def export_rows(records):
    for record in records:
        yield {**record["fields"], "source_file": record["source_file"], "source_refs": "\n".join(record["sources"]),
               "match_status": record["match_status"],
               "matching_pids": "; ".join(m["pid"] for m in record["matches"]),
               "match_reasons": "\n".join(m["pid"] + ": " + "; ".join(m["reasons"]) for m in record["matches"]),
               "review_notes": "\n".join(record["warnings"] + record.get("match_notes", []))}


def export_xlsx(records):
    book = Workbook()
    sheet = book.active
    sheet.title = "Extracted participants"
    sheet.append(list(EXPORT_COLUMNS.values()))
    for row in export_rows(records):
        values = [str(row.get(key, "")) for key in EXPORT_COLUMNS]
        if any(len(value) > 32767 for value in values):
            raise WordExportError("A value exceeds Excel's 32,767-character cell limit. Download CSV to preserve the full text.")
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            # Explicit text cells retain leading zeros and never become Excel formulas.
            cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, key in enumerate(EXPORT_COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 45 if key in ("bio_short", "position", "review_notes", "match_reasons", "unmapped_text") else 23
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def export_csv(records):
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(EXPORT_COLUMNS.values())
    for row in export_rows(records):
        values = [str(row.get(key, "")) for key in EXPORT_COLUMNS]
        writer.writerow(["'" + value if value.lstrip().startswith(("=", "+", "-", "@")) else value for value in values])
    return output.getvalue().encode("utf-8-sig")
