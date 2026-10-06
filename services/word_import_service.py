"""Read Word extraction exports and reuse the Master Tracker participant review."""

import csv
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO, StringIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from domain.models.participant import Participant
from domain.reporting import normalize_text
from services.import_validation import EDITABLE_FIELDS, SNAPSHOT_FIELDS, participant_errors, partial_snapshot_errors
from services.imports.participant_review import PROFILE_FIELDS, ReviewMatchError, annotate_participant_reviews, find_returning_participant
from services.word_export_service import EXPORT_COLUMNS
from services.word_extraction_service import FIELDS, MAX_FILE_BYTES, MAX_RECORDS, name_key, parse_date


class WordImportError(ValueError):
    pass


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    return str(value).strip()


def _table_rows(content, filename):
    if len(content) > MAX_FILE_BYTES:
        raise WordImportError("The extracted file exceeds 10 MB.")
    if filename.lower().endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise WordImportError("Save the CSV as UTF-8 before uploading it.") from exc
        try:
            for row in csv.reader(StringIO(text, newline=""), strict=True):
                # Reverse the formula protection applied by our CSV exporter.
                yield [value[1:] if value.startswith("'") and value[1:].lstrip().startswith(("=", "+", "-", "@")) else value for value in row]
        except csv.Error as exc:
            raise WordImportError("The CSV is unreadable. Re-save it as UTF-8 CSV.") from exc
    elif filename.lower().endswith(".xlsx"):
        book = None
        try:
            with ZipFile(BytesIO(content)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                    raise WordImportError("The workbook's expanded contents exceed 50 MB.")
            book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
            sheet = book["Extracted participants"] if "Extracted participants" in book.sheetnames else book.worksheets[0]
            for row in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise WordImportError("The workbook contains formulas. Paste their values before importing.")
                yield [cell.value for cell in row]
        except (BadZipFile, KeyError, ValueError, IndexError) as exc:
            if isinstance(exc, WordImportError):
                raise
            raise WordImportError("This is not a readable extracted Excel workbook.") from exc
        finally:
            if book:
                book.close()
    else:
        raise WordImportError("Upload an extracted .xlsx or UTF-8 .csv file.")


def read_export(content, filename):
    aliases = {normalize_text(label): key for key, label in EXPORT_COLUMNS.items()}
    aliases.update({normalize_text(key): key for key in EXPORT_COLUMNS})
    rows = iter(_table_rows(content, filename))
    header = next(rows, [])
    columns = [aliases.get(normalize_text(_text(value))) for value in header]
    known = [column for column in columns if column]
    if "name" not in known or not {"representing_country", "country_label"}.intersection(known):
        raise WordImportError("Use the extracted table with Name and Country CID or Country as stated columns.")
    if len(known) != len(set(known)):
        raise WordImportError("The table has duplicate columns. Keep one column for each field.")
    records = []
    for number, row in enumerate(rows, 2):
        if not any(_text(value) for value in row):
            continue
        if len(row) > len(header) and any(_text(value) for value in row[len(header):]):
            raise WordImportError(f"Row {number} has more values than column headings.")
        if len(records) >= MAX_RECORDS:
            raise WordImportError("Use at most 1,000 participant rows per import.")
        values = {key: _text(value) for key, value in zip(columns, row) if key}
        fields = {key: values.get(key, "") for key in FIELDS}
        source = {"file": filename, "row": number, "fields": fields,
                  "source_file": values.get("source_file", ""), "source_refs": values.get("source_refs", ""),
                  "extra_columns": [{"label": _text(label), "value": _text(value)}
                                    for key, label, value in zip(columns, header, row) if not key and _text(value)]}
        records.append({"fields": fields, "source": source})
    if not records:
        raise WordImportError("No participant rows were found in the extracted file.")
    return records


def _country(value, context):
    text = _text(value)
    return (text if text in context.countries else context.resolve_country(text)) or text


def convert_record(record, context):
    fields = record["fields"]
    participant = {key: value for key, value in fields.items() if key in EDITABLE_FIELDS and value}
    participant["representing_country"] = _country(fields["representing_country"] or fields["country_label"], context)
    participant.setdefault("name", "")
    participant.setdefault("pob", None)
    participant["birth_country"] = _country(fields["birth_country"], context) or None
    if fields["citizenships"]:
        # A complete catalog label may itself contain a comma.
        whole = context.resolve_country(fields["citizenships"])
        participant["citizenships"] = [whole] if whole else [_country(v, context) for v in fields["citizenships"].replace(",", ";").split(";") if v.strip()]
    for key in ("dob", "travel_doc_issue_date", "travel_doc_expiry_date"):
        if participant.get(key):
            participant[key] = parse_date(participant[key]) or participant[key]
    gender = normalize_text(participant.get("gender", ""))
    if gender in ("m", "male", "f", "female"):
        participant["gender"] = "Male" if gender in ("m", "male") else "Female"
    grade = normalize_text(participant.get("grade", ""))
    if grade in ("normal", "excellent", "black list", "blacklist"):
        participant["grade"] = "1" if grade == "normal" else "2" if grade == "excellent" else "0"
    # Explicit transport labels are normalized; no missing travel data is inferred.
    transport = normalize_text(participant.get("transportation", ""))
    if transport in ("air", "airplane", "flight"):
        participant["transportation"] = "Air (Airplane)"
    participant["_word_source"] = record["source"]
    participant["_source_event"] = fields["event_reference"]
    return participant


def review_records(participants, repo, context):
    """Resolve fresh candidates, then use the shared profile comparison rules."""
    result = []
    by_pid = {person["pid"]: person for person in context.participants}
    for source in participants:
        item = dict(source)
        for key in ("_changes", "_match_error", "_match_candidates", "_field_errors"):
            item.pop(key, None)
        if item.get("_review"):
            try:
                item = annotate_participant_reviews([item], repo)[0]
            except ReviewMatchError as exc:
                item["_match_error"] = str(exc)
            item["_match_choice"] = item["_review"]["pid"]
            result.append(item)
            continue
        probe = {"fields": {key: "" for key in FIELDS} | {key: _text(value) for key, value in item.items() if key in FIELDS}}
        context.check(probe)
        candidates = probe.get("matches", [])
        eligible = [candidate for candidate in candidates if
                    name_key(candidate["name"]) == name_key(item.get("name")) and
                    by_pid[candidate["pid"]]["country"] == item.get("representing_country") and
                    (not parse_date(item.get("dob")) or not candidate["dob"] or candidate["dob"] == parse_date(item.get("dob")))]
        choice = item.get("_match_choice", "")
        if choice == "new" and eligible:
            item["_match_error"] = "An existing participant matches this identity. Choose that PID instead of creating a duplicate."
        elif choice and choice != "new" and choice not in {candidate["pid"] for candidate in candidates}:
            item["_match_error"] = "The selected participant is no longer a candidate. Review the identity again."
        else:
            chosen = choice if choice and choice != "new" else eligible[0]["pid"] if len(eligible) == 1 and not choice else None
            if chosen:
                person = by_pid[chosen]
                item["_review"] = {"pid": chosen, "identity": {"name": person["name"], "dob": person["dob"], "representing_country": person["country"]}}
                item = annotate_participant_reviews([item], repo)[0]
                item["_match_choice"] = chosen
            elif candidates and choice != "new":
                item["_match_error"] = "Review the possible existing participants and select a PID or confirm a new participant."
        item["_match_candidates"] = candidates
        result.append(item)
    repeated = defaultdict(list)
    for item in result:
        item.pop("_duplicate_note", None)
        dob = parse_date(item.get("dob"))
        if dob and item.get("representing_country"):
            repeated[(name_key(item.get("name")), dob, item["representing_country"])].append(item)
    for group in repeated.values():
        if len(group) > 1:
            for item in group:
                item["_duplicate_note"] = "This name, DOB, and country also appear on another row. Selected copies will share one participant PID."
    return result


def row_errors(item, repo, context):
    existing = find_returning_participant(item, repo) if item.get("_review") else None
    if existing:
        selected = set(item["_review"].get("accepted_fields", [])) & PROFILE_FIELDS
        payload = existing.model_dump() | {key: item[key] for key in selected if key in item}
    else:
        selected = PROFILE_FIELDS
        payload = {key: value for key, value in item.items() if key in Participant.model_fields}
    errors = participant_errors(payload, allow_missing_dob=bool(existing))
    for key in ("representing_country", "birth_country"):
        if key in selected and payload.get(key) and payload[key] not in context.countries:
            errors[key] = "Enter a country CID from the country catalog."
    if "citizenships" in selected and payload.get("citizenships") and any(value not in context.countries for value in payload["citizenships"]):
        errors["citizenships"] = "Use country CIDs separated by semicolons."
    if payload.get("dob") and not parse_date(payload["dob"]):
        errors["dob"] = "Enter a valid, unambiguous DOB as YYYY-MM-DD."
    errors.update(partial_snapshot_errors(item))
    for key in ("travel_doc_issue_date", "travel_doc_expiry_date"):
        if item.get(key) and not parse_date(item[key]):
            errors[key] = "Enter a valid, unambiguous date as YYYY-MM-DD."
    if item.get("_match_error"):
        errors["name"] = item["_match_error"]
    return errors
