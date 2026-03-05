from __future__ import annotations

import re
from datetime import datetime
from typing import Any

from docx import Document

_LABEL_ALIASES = {
    "country": "representing_country",
    "name & last name": "name",
    "name": "name",
    "name last name": "name",
    "dob": "dob",
    "date of birth": "dob",
    "pob": "pob",
    "place of birth": "pob",
    "passport number": "travel_doc_number",
    "gender": "gender",
    "diet restrictions": "diet_restrictions",
    "institution": "organization",
    "institution/organization": "organization",
    "organization": "organization",
    "position": "position",
    "rank": "rank",
    "phone": "phone",
    "email": "email",
    "transportation": "transportation",
    "arrival": "arrival",
    "departure": "departure",
    "short professional biography": "bio_short",
}

_DATE_FORMATS = [
    "%B/%d/%Y",
    "%b/%d/%Y",
    "%m/%d/%Y",
    "%m/%d/%y",
    "%d/%m/%Y",
    "%d.%m.%Y",
    "%d. %m. %Y",
]


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").replace("\xa0", " ")).strip()


def _normalize_label(label: str) -> str:
    normalized = re.sub(r"\s+", " ", label.lower()).strip(" :")
    return _LABEL_ALIASES.get(normalized, normalized.replace(" ", "_"))


def _normalize_date(raw: str) -> str:
    value = _clean(raw).strip("./")
    if not value:
        return ""

    value = value.replace("-", "/")
    value = re.sub(r"\s+", "", value)

    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue

    # Serbian/Croatian style: 13. 2. 1965.
    m = re.match(r"^(\d{1,2})\.?[./](\d{1,2})\.?[./](\d{2,4})\.?$", value)
    if m:
        day, month, year = m.groups()
        year = year if len(year) == 4 else f"19{year}"
        try:
            return datetime(int(year), int(month), int(day)).strftime("%Y-%m-%d")
        except ValueError:
            return ""
    return ""


def _normalize_gender(raw: str) -> str:
    text = _clean(raw).lower()
    if text in {"f", "female", "famele"}:
        return "Female"
    if text in {"m", "male"}:
        return "Male"
    return ""


def _normalize_bool(raw: str) -> bool | str:
    text = _clean(raw).lower()
    if text in {"no", "n"}:
        return False
    if text in {"yes", "y"}:
        return True
    return ""


def _extract_pattern_1(doc: Document) -> list[dict[str, Any]]:
    participants: list[dict[str, Any]] = []
    for table in doc.tables:
        fields: dict[str, Any] = {}
        for row in table.rows:
            cells = [_clean(cell.text) for cell in row.cells if _clean(cell.text)]
            if len(cells) < 2:
                continue
            i = 0
            while i + 1 < len(cells):
                key = _normalize_label(cells[i])
                fields[key] = cells[i + 1]
                i += 2

        if fields.get("name"):
            participants.append(fields)
    return participants


def _extract_pattern_2(lines: list[str]) -> list[dict[str, Any]]:
    participants: list[dict[str, Any]] = []
    date_pattern = re.compile(r"(\d{1,2}[/.]\d{1,2}[/.]\d{2,4})")

    for line in lines:
        if not date_pattern.search(line):
            continue
        parts = re.split(r"\t+|\s{2,}", line)
        if len(parts) < 3:
            continue

        name = _clean(parts[0])
        dob_match = date_pattern.search(line)
        if not name or not dob_match:
            continue

        rank = _clean(parts[2]) if len(parts) > 2 else ""
        tail = _clean(parts[3]) if len(parts) > 3 else ""
        position, _, organization = tail.partition(",")
        if "/" in position and not organization:
            position, _, organization = position.partition("/")

        participants.append(
            {
                "name": name,
                "dob": dob_match.group(1),
                "rank": rank,
                "position": _clean(position),
                "organization": _clean(organization),
            }
        )

    return participants


def _extract_pattern_3(lines: list[str]) -> list[dict[str, Any]]:
    participants: list[dict[str, Any]] = []
    country = ""
    country_header = re.compile(r"^\d+\.\s+(.+)$")
    list_item = re.compile(r"^\d+\.\s+([A-Za-zČĆŽŠĐčćžšđ\-\s]+)(?:\(|$)")

    for line in lines:
        header_match = country_header.match(line)
        has_date = bool(re.search(r"\d{1,2}[./]\s*\d{1,2}[./]", line))
        if header_match and len(header_match.group(1).split()) <= 4 and not has_date:
            country = _clean(header_match.group(1))
            continue

        m = list_item.match(line)
        if not m:
            continue

        name = _clean(m.group(1))
        if not name:
            continue

        dob_match = re.search(r"\(([^)]+)\)", line)
        participants.append(
            {
                "name": name,
                "dob": dob_match.group(1) if dob_match else "",
                "representing_country": country,
            }
        )
    return participants


def extract_docx_text(path: str, *, max_length: int = 12000) -> str:
    """Extract plain text from DOCX paragraphs and tables for LLM parsing."""

    doc = Document(path)
    lines: list[str] = []

    for paragraph in doc.paragraphs:
        text = _clean(paragraph.text)
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = _clean(cell.text)
                if text:
                    lines.append(text)

    return "\n".join(lines)[:max_length]


def parse_docx(path: str) -> list[dict[str, Any]]:
    """Parse a DOCX file and return best-effort participant dictionaries."""

    doc = Document(path)
    lines = [_clean(p.text) for p in doc.paragraphs if _clean(p.text)]

    participants = _extract_pattern_1(doc)
    if not participants:
        participants = _extract_pattern_2(lines)
    if not participants:
        participants = _extract_pattern_3(lines)

    normalized: list[dict[str, Any]] = []
    for participant in participants:
        item = dict(participant)
        if "dob" in item:
            item["dob"] = _normalize_date(str(item.get("dob", "")))
        if "arrival" in item:
            item["arrival"] = _normalize_date(str(item.get("arrival", "")))
        if "departure" in item:
            item["departure"] = _normalize_date(str(item.get("departure", "")))
        if "gender" in item:
            item["gender"] = _normalize_gender(str(item.get("gender", "")))
        if "diet_restrictions" in item:
            maybe_bool = _normalize_bool(str(item.get("diet_restrictions", "")))
            if maybe_bool != "":
                item["diet_restrictions"] = maybe_bool
        normalized.append(item)

    return normalized
