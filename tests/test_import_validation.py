import json
from io import BytesIO
from pathlib import Path

import pytest
from flask import Flask
from openpyxl import Workbook, load_workbook

from domain.models.participant import Participant
from services.import_validation import event_errors, participant_errors, snapshot_errors
from services.imports.participant_review import annotate_participant_reviews
from services.upload_service import UploadError, upload_preview_data
from tests.test_upload_service import FakeEventRepo, FakeParticipantRepo, FakeParticipantEventRepo, _base_event, _base_participant
from tests.test_upload_service import isolated_upload_session
from utils.costs import parse_cost, read_grand_total


@pytest.mark.parametrize("amount", [64711.19, "64,711.19", "64.711,19", "64 711,19 EUR"])
def test_total_comes_from_label_instead_of_line_item(tmp_path, amount):
    from services.import_service_v2 import _read_event_header_block, validate_excel_file_for_import
    from tests.test_import_service_gender import _workbook_bytes_with_gender
    wb = load_workbook(BytesIO(_workbook_bytes_with_gender("Male")))
    ws = wb["COST Overview"]
    ws["B15"] = 440.32
    ws.merge_cells("C28:D28")
    ws["C28"] = " gRaNd   TOTAL "
    ws["E28"] = amount
    path = tmp_path / "PFE26M1.xlsx"
    wb.save(path)
    assert _read_event_header_block(str(path))[-1] == 64711.19
    assert validate_excel_file_for_import(str(path))[0]


def test_zero_total_is_valid():
    ws = Workbook().active
    ws.append(["GRAND TOTAL:", 0])
    assert read_grand_total(ws) == 0


def test_unlabelled_line_item_does_not_become_event_cost():
    ws = Workbook().active
    ws["B15"] = 440.32
    assert read_grand_total(ws) is None


@pytest.mark.parametrize("values", [[None], ["=SUM(B1:B10)"], [10, 20], [-20], [float("inf")]])
def test_missing_invalid_or_ambiguous_totals_are_rejected(values):
    ws = Workbook().active
    ws.append(["GRAND TOTAL", *values])
    with pytest.raises(ValueError):
        read_grand_total(ws)


@pytest.mark.parametrize("value", ["NaN", "Infinity", "-1", True, "abc", "64,71.19", "64,711"])
def test_invalid_cost_is_not_silently_discarded(value):
    with pytest.raises(ValueError):
        parse_cost(value)


def test_email_case_does_not_change_profile():
    repo = FakeParticipantRepo()
    person = Participant.model_validate({**_base_participant(), "pid": "P1234", "email": " Jane@EXAMPLE.COM "})
    repo.participants[person.pid] = person
    assert person.email == "jane@example.com"
    reviewed = annotate_participant_reviews([_base_participant(email="JANE@example.com")], repo)[0]
    assert "email" not in reviewed["_changes"]


def test_all_invalid_fields_are_reported():
    errors = participant_errors(_base_participant(name="  ", gender="invalid", grade=8, email="invalid", phone="invalid", dob=None))
    assert {"name", "gender", "grade", "email", "phone", "dob"} <= errors.keys()
    errors = snapshot_errors(_base_participant(traveling_from=" ", returning_to="", travel_doc_type="bad", iban_type="bad", travel_doc_issue_date="2034-01-01", travel_doc_expiry_date="2024-01-01"))
    assert {"traveling_from", "returning_to", "travel_doc_type", "iban_type", "travel_doc_issue_date", "travel_doc_expiry_date"} <= errors.keys()
    errors = event_errors({**_base_event(), "cost": "abc", "start_date": "bad", "type": "bad"})
    assert {"cost", "start_date", "type"} <= errors.keys()


@pytest.mark.parametrize("bad_field,bad_value", [("iban_type", "bad"), ("returning_to", " "), ("email", "invalid"), ("intl_authority", "maybe")])
def test_invalid_second_participant_rejected_before_transaction(monkeypatch, bad_field, bad_value):
    import services.upload_service as uploads
    monkeypatch.setattr(uploads.mongodb, "start_session", lambda: pytest.fail("No transaction should begin"), raising=False)
    with pytest.raises(UploadError, match=bad_field.replace("_", " ").capitalize()):
        upload_preview_data(
            {"event": _base_event(), "participants": [_base_participant(), _base_participant(name="Second Person", **{bad_field: bad_value})]},
            event_repo=FakeEventRepo(), participant_repo=FakeParticipantRepo(), participant_event_repo=FakeParticipantEventRepo(),
        )


def test_invalid_event_cost_rejected_before_transaction(monkeypatch):
    import services.upload_service as uploads
    monkeypatch.setattr(uploads.mongodb, "start_session", lambda: pytest.fail("No transaction should begin"), raising=False)
    with pytest.raises(UploadError, match="Cost"):
        upload_preview_data(
            {"event": {**_base_event(), "cost": "bad"}, "participants": []},
            event_repo=FakeEventRepo(), participant_repo=FakeParticipantRepo(), participant_event_repo=FakeParticipantEventRepo(),
        )


def test_valid_email_and_document_number_are_saved(isolated_upload_session):
    snapshots = FakeParticipantEventRepo()
    result = upload_preview_data(
        {"event": _base_event(), "participants": [_base_participant(email="JANE@EXAMPLE.COM", travel_doc_number="AB123456")]},
        event_repo=FakeEventRepo(), participant_repo=FakeParticipantRepo(), participant_event_repo=snapshots,
    )
    assert result["participants"][0].email == "jane@example.com"
    assert snapshots.snapshots[0].travel_doc_number == "AB123456"


def test_future_dob_and_invalid_country_lists_are_rejected():
    assert "dob" in participant_errors(_base_participant(dob="2999-01-01"))
    assert "citizenships" in participant_errors(_base_participant(citizenships=[12]))


def test_missing_total_is_reported_during_workbook_validation(tmp_path):
    from services.import_service_v2 import validate_excel_file_for_import
    from tests.test_import_service_gender import _workbook_bytes_with_gender
    path = tmp_path / "missing-total.xlsx"
    path.write_bytes(_workbook_bytes_with_gender("Male"))
    ok, errors, _ = validate_excel_file_for_import(str(path))
    assert not ok
    assert any("GRAND TOTAL" in message for message in errors)


def test_preview_highlights_errors_preserves_bad_input_and_accepts_missing_fields(tmp_path, monkeypatch):
    import routes.imports as routes
    from utils.participants import initialize_cache
    initialize_cache(None)
    monkeypatch.setattr(routes, "ParticipantRepository", FakeParticipantRepo)
    app = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / "templates"))
    app.config.update(TESTING=True, LOGIN_DISABLED=True, UPLOADS_DIR=str(tmp_path), SECRET_KEY="test")
    app.register_blueprint(routes.imports_bp)
    for endpoint in ("main.show_home", "participants.show_participants", "events.show_events", "auth.login"):
        app.add_url_rule(f"/test/{endpoint}", endpoint=endpoint, view_func=lambda: "")
    path = tmp_path / "fields.preview.json"
    path.write_text(json.dumps({"event": {**_base_event(), "cost": 10}, "participants": [_base_participant(email="invalid", intl_authority=False)]}))
    with app.test_client() as client:
        html = client.get("/imports/preview/fields.preview.json").get_data(as_text=True)
        assert 'is-invalid" id="participant_0_email"' in html
        client.post("/imports/preview/fields.preview.json", data={"event[cost]": "bad", "participants[0][intl_authority]": "maybe", "participants[0][organization]": "New organization"})
        saved = json.loads(path.read_text())
        assert saved["event"]["cost"] == "bad"
        assert saved["participants"][0]["intl_authority"] == "maybe"
        assert saved["participants"][0]["organization"] == "New organization"
        html = client.get("/imports/preview/fields.preview.json").get_data(as_text=True)
        assert 'is-invalid" id="event_cost"' in html
        assert 'is-invalid" id="participant_0_intl_authority"' in html
