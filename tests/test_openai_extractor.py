from __future__ import annotations

from utils.docx_parser import extract_docx_text
from utils.openai_extractor import extract_participants
from docx import Document


class _FakeResponses:
    def __init__(self, output_text: str) -> None:
        self._output_text = output_text

    def create(self, **_kwargs):
        class _Resp:
            def __init__(self, output_text: str) -> None:
                self.output_text = output_text

        return _Resp(self._output_text)


class _FakeClient:
    def __init__(self, output_text: str) -> None:
        self.responses = _FakeResponses(output_text)


def test_openai_extractor_parses_json_fence() -> None:
    client = _FakeClient(
        """```json
{"participants":[{"name":"Jane Doe","citizenships":"C001"}]}
```"""
    )

    participants = extract_participants("sample text", client=client)

    assert len(participants) == 1
    assert participants[0]["name"] == "Jane Doe"
    assert participants[0]["citizenships"] == ["C001"]


def test_extract_docx_text_reads_paragraphs_and_table(tmp_path) -> None:
    file_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Header")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "NAME"
    table.cell(0, 1).text = "John Doe"
    doc.save(file_path)

    text = extract_docx_text(str(file_path))

    assert "Header" in text
    assert "NAME" in text
    assert "John Doe" in text
