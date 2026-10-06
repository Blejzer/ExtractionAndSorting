from contextlib import nullcontext
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
from pathlib import Path

from flask import Flask
from openpyxl import Workbook
import pytest
from pymongo.errors import ServerSelectionTimeoutError

from domain.models.event import Event
from domain.models.participant import Participant
from repositories.event_repository import EventRepository
from repositories.participant_event_repository import ParticipantEventRepository
from routes import word_extraction, word_import
from services import upload_service
from services.imports.participant_review import ReviewMatchError, find_returning_participant
from services.word_export_service import EXPORT_COLUMNS, export_csv, export_rows
from services.word_import_service import WordImportError, convert_record, read_export, review_records
from services.word_matching_service import WordMatchContext
from tests.test_returning_participant_review import ReturningRepo
from tests.test_upload_service import FakeEventRepo, _base_event, _base_participant
from tests.test_word_extraction import COUNTRIES, login, record


class People(ReturningRepo):
    def find_by_pid(self, pid):
        return self.participants.get(pid)


class Events(FakeEventRepo):
    def find_all(self):
        return list(self.events.values())

    def add_participants(self, eid, participant_ids, *, session=None):
        event = self.events.get(eid)
        if event:
            event.participants = list(dict.fromkeys(event.participants + participant_ids))
        return event


class Snapshots:
    def __init__(self):
        self.rows = {}
        self.writes = 0

    def upsert_partial(self, pid, eid, fields, sources, *, session=None):
        self.writes += 1
        row = self.rows.setdefault((pid, eid), {"participant_id": pid, "event_id": eid})
        row.update(fields)
        for source in sources:
            evidence = row.setdefault("word_import_sources", [])
            if source not in evidence:
                evidence.append(source)


def exported(fmt="xlsx", **fields):
    item = record()
    item["fields"].update({"gender": "Female", "event_reference": "EVT-001", **fields})
    if fmt == "csv":
        return export_csv([item])
    return legacy_excel([item])


def legacy_excel(records):
    # Existing Excel tables remain valid import inputs; exports now use CSV.
    book = Workbook()
    sheet = book.active
    sheet.append(list(EXPORT_COLUMNS.values()))
    for row in export_rows(records):
        sheet.append([str(row.get(key, "")) for key in EXPORT_COLUMNS])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


@pytest.fixture
def setup(tmp_path, monkeypatch):
    people, events, snapshots = People(), Events(), Snapshots()
    events.save(Event(**_base_event(), participants=["POLD"], training_areas=["cybercrime"]))
    app = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / "templates"))
    app.config.update(TESTING=True, SECRET_KEY="import-test", WORD_EXTRACTION_DIR=str(tmp_path))
    for endpoint in ("main.show_home", "participants.show_participants", "events.show_events", "imports.upload_form", "auth.login", "auth.logout"):
        app.add_url_rule("/" + endpoint, endpoint=endpoint, view_func=lambda: "stub")
    app.add_url_rule("/participants/<pid>", endpoint="participants.participant_detail", view_func=lambda pid: pid)
    app.add_url_rule("/events/<eid>", endpoint="events.event_detail", view_func=lambda eid: eid)
    app.register_blueprint(word_extraction.word_extraction_bp)
    app.register_blueprint(word_import.word_import_bp)
    monkeypatch.setattr(word_import, "EventRepository", lambda: events)
    monkeypatch.setattr(word_import, "ParticipantRepository", lambda: people)
    context = lambda: WordMatchContext([person.model_dump() for person in people.participants.values()], COUNTRIES)
    monkeypatch.setattr(word_import, "load_match_context", context)
    monkeypatch.setattr(word_extraction, "load_match_context", context)
    monkeypatch.setattr(upload_service, "ParticipantEventRepository", lambda: snapshots)
    monkeypatch.setattr(upload_service, "refresh_participant_cache", lambda: None)
    class Session:
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def start_transaction(self):
            return nullcontext()
    monkeypatch.setattr(upload_service.mongodb, "start_session", Session, raising=False)
    return app, people, events, snapshots


def returning(people, **overrides):
    data = _base_participant(pid="P0814", name="Ana TEST", representing_country="AL_TEST", birth_country="AL_TEST", citizenships=["AL_TEST"], dob="1980-03-25", phone="+385999999", organization="Stored organization")
    data.update(overrides)
    person = Participant.model_validate(data, context={"allow_missing_dob": True})
    people.participants[person.pid] = person
    return person


def start(client, token, content=None, fmt="xlsx", event_id="EVT-001", **form):
    return client.post("/imports/word/import", data={"csrf_token": token, "event_id": event_id, "file": (BytesIO(content or exported(fmt)), f"extracted.{fmt}"), **form}, content_type="multipart/form-data")


def bundle(app, location):
    return json.loads((Path(app.config["WORD_EXTRACTION_DIR"]) / (location.rsplit("/", 1)[-1] + ".json")).read_text())["data"]


def submit(client, token, app, location, *, accepted=None, **edits):
    data = {"csrf_token": token, "import_now": "1"}
    for index, row in enumerate(bundle(app, location)["participants"]):
        if row["_include"]:
            data[f"include[{index}]"] = "1"
        for field in (row.get("_review") or {}).get("accepted_fields", []) if accepted is None else accepted:
            data[f"accept[{index}][{field}]"] = "1"
    return client.post(location, data=data | edits, follow_redirects=True)


@pytest.mark.parametrize("fmt", ["xlsx", "csv"])
def test_export_round_trip_preserves_country_phone_documents_and_bio(fmt):
    data = exported(fmt, phone="+385111111", travel_doc_number="001234", bio_short="- First line\nSecond line", service_number="0099")
    rows = read_export(data, f"export.{fmt}")
    context = WordMatchContext([], COUNTRIES)
    person = convert_record(rows[0], context)
    assert person["phone"] == "+385111111"
    assert person["travel_doc_number"] == "001234"
    assert person["bio_short"] == "- First line\nSecond line"
    assert person["representing_country"] == "AL_TEST"
    assert person["_word_source"]["fields"]["service_number"] == "0099"


@pytest.mark.parametrize("fmt", ["xlsx", "csv"])
def test_preview_is_read_only_then_adds_new_participant_to_existing_event(setup, fmt):
    app, people, events, snapshots = setup
    client = app.test_client()
    token = login(client)
    location = start(client, token, fmt=fmt).headers["Location"]
    page = client.get(location)
    assert page.status_code == 200
    assert b"New participant" in page.data
    assert not people.participants and not snapshots.rows
    page = submit(client, token, app, location)
    assert b"Participants imported" in page.data
    assert list(people.participants) == ["P0001"]
    assert events.events["EVT-001"].participants == ["POLD", "P0001"]
    assert events.events["EVT-001"].training_areas == ["cybercrime"]
    assert ("P0001", "EVT-001") in snapshots.rows
    assert "transportation" not in snapshots.rows[("P0001", "EVT-001")]
    assert snapshots.rows[("P0001", "EVT-001")]["word_import_sources"][0]["fields"]["name"] == "Ana TEST"
    writes = snapshots.writes
    submit(client, token, app, location)
    assert people.counter == 2 and snapshots.writes == writes


def test_returning_participant_is_marked_and_only_selected_profile_changes_apply(setup):
    app, people, events, snapshots = setup
    old = returning(people)
    client = app.test_client()
    token = login(client)
    location = start(client, token, exported(phone="+385111111", organization="File organization", travel_doc_number="001234")).headers["Location"]
    page = client.get(location)
    assert b"Returning participant: P0814" in page.data
    assert b"bg-warning-subtle" in page.data and b"Stored organization" in page.data
    assert people.participants[old.pid].phone == old.phone
    page = submit(client, token, app, location, accepted=["phone"])
    assert b"Participants imported" in page.data
    assert list(people.participants) == ["P0814"] and people.counter == 1
    assert people.participants[old.pid].phone == "+385111111"
    assert people.participants[old.pid].organization == "Stored organization"
    assert people.participants[old.pid].birth_country == "AL_TEST"
    assert events.events["EVT-001"].participants == ["POLD", "P0814"]
    assert snapshots.rows[(old.pid, "EVT-001")]["travel_doc_number"] == "001234"


def test_missing_stored_dob_keeps_pid_and_offers_file_value(setup):
    app, people, _, _ = setup
    returning(people, dob=None)
    client = app.test_client()
    token = login(client)
    location = start(client, token).headers["Location"]
    assert "dob" in bundle(app, location)["participants"][0]["_review"]["accepted_fields"]
    assert b"Participants imported" in submit(client, token, app, location).data
    assert people.participants["P0814"].dob.date().isoformat() == "1980-03-25"
    assert people.counter == 1


@pytest.mark.parametrize("dob,expected", [("1967/2/10", "1967-02-10"), ("1967-2-10", "1967-02-10"), ("1971/11/26", "1971-11-26"), ("1971-11-26", "1971-11-26"), ("11/26/1971", "1971-11-26")])
def test_edited_year_first_dates_are_normalized_before_database_import(setup, dob, expected):
    app, people, _, snapshots = setup
    client = app.test_client()
    token = login(client)
    location = start(client, token).headers["Location"]
    page = submit(client, token, app, location, **{
        "participants[0][dob]": dob,
        "participants[0][travel_doc_issue_date]": "2020/2/10",
        "participants[0][travel_doc_expiry_date]": "2030-2-10",
    })
    assert b"Participants imported" in page.data
    assert people.participants["P0001"].dob.date().isoformat() == expected
    snapshot = snapshots.rows[("P0001", "EVT-001")]
    assert snapshot["travel_doc_issue_date"].date().isoformat() == "2020-02-10"
    assert snapshot["travel_doc_expiry_date"].date().isoformat() == "2030-02-10"


def test_partial_travel_updates_keep_existing_snapshot_and_other_events(setup):
    app, people, events, snapshots = setup
    person = returning(people)
    current = {"participant_id": person.pid, "event_id": "EVT-001", "transportation": "Air (Airplane)", "bank_name": "Keep this bank", "travel_doc_number": "old"}
    previous = {"participant_id": person.pid, "event_id": "OLD", "travel_doc_number": "previous"}
    snapshots.rows[(person.pid, "EVT-001")] = dict(current)
    snapshots.rows[(person.pid, "OLD")] = dict(previous)
    client = app.test_client()
    token = login(client)
    location = start(client, token, exported(travel_doc_number="001234")).headers["Location"]
    assert b"Participants imported" in submit(client, token, app, location).data
    updated = snapshots.rows[(person.pid, "EVT-001")]
    assert updated["transportation"] == current["transportation"]
    assert updated["bank_name"] == current["bank_name"]
    assert updated["travel_doc_number"] == "001234"
    assert snapshots.rows[(person.pid, "OLD")] == previous


@pytest.mark.parametrize("bad", [{"gender": ""}, {"dob": ""}, {"dob": "2/10/1980"}, {"representing_country": "INVALID"}, {"travel_doc_type": "bad"}, {"travel_doc_issue_date": "2034-01-01", "travel_doc_expiry_date": "2024-01-01"}, {"iban_type": "bad"}])
def test_invalid_selected_fields_block_all_writes(setup, bad):
    app, people, events, snapshots = setup
    client = app.test_client()
    token = login(client)
    location = start(client, token, exported(**bad)).headers["Location"]
    page = submit(client, token, app, location)
    assert b"Participants imported" not in page.data
    assert not people.participants and people.counter == 1 and not snapshots.rows
    assert events.events["EVT-001"].participants == ["POLD"]


def test_unselected_invalid_profile_value_does_not_block_returning_participant(setup):
    app, people, _, _ = setup
    old = returning(people)
    client = app.test_client()
    token = login(client)
    location = start(client, token, exported(email="invalid-email")).headers["Location"]
    assert b"Participants imported" in submit(client, token, app, location, accepted=[]).data
    assert people.participants[old.pid].email == old.email


def test_unmodified_legacy_birth_country_does_not_block_attendance_import(setup):
    app, people, _, _ = setup
    old = returning(people, birth_country="NA")
    client = app.test_client()
    token = login(client)
    location = start(client, token).headers["Location"]
    assert b"Participants imported" in submit(client, token, app, location, accepted=[]).data
    assert people.participants[old.pid].birth_country == "NA"


def test_mixed_event_references_default_to_only_target_event_rows(setup):
    app, people, events, snapshots = setup
    first, second = record(), record(name="Bora TEST", dob="1981-03-26")
    first["fields"].update(gender="Female", event_reference="EVT-001")
    second["fields"].update(event_reference="OTHER")  # Invalid unselected rows must not block import.
    client = app.test_client()
    token = login(client)
    location = start(client, token, legacy_excel([first, second])).headers["Location"]
    data = bundle(app, location)
    assert [row["_include"] for row in data["participants"]] == [True, False]
    assert b"Participants imported" in submit(client, token, app, location).data
    assert len(people.participants) == 1 and len(snapshots.rows) == 1


def test_repeated_new_identity_uses_one_pid_and_keeps_fields_absent_from_later_copy(setup):
    app, people, events, snapshots = setup
    first, second = record(), record(name="TEST, Ana")
    first["fields"].update(gender="Female", grade="0", birth_country="Albania", event_reference="EVT-001")
    second["fields"].update(gender="Female", position="Specialist", event_reference="EVT-001")
    client = app.test_client()
    token = login(client)
    location = start(client, token, legacy_excel([first, second])).headers["Location"]
    assert b"share one participant PID" in client.get(location).data
    assert b"Participants imported" in submit(client, token, app, location).data
    assert people.counter == 2 and len(people.participants) == 1
    assert people.participants["P0001"].grade == 0
    assert people.participants["P0001"].birth_country == "AL_TEST"
    assert people.participants["P0001"].position == "Specialist"
    assert events.events["EVT-001"].participants == ["POLD", "P0001"]
    assert len(snapshots.rows[("P0001", "EVT-001")]["word_import_sources"]) == 2


def test_can_create_new_event_after_completing_event_details(setup):
    app, people, events, _ = setup
    client = app.test_client()
    token = login(client)
    location = start(client, token, exported(event_reference="NEW"), event_id="__new__").headers["Location"]
    assert b"Participants imported" not in submit(client, token, app, location).data
    assert not people.participants
    page = submit(client, token, app, location, **{"event[title]": "New training", "event[start_date]": "2026-10-06", "event[end_date]": "2026-10-09", "event[place]": "Test city"})
    assert b"Participants imported" in page.data
    assert events.events["NEW"].participants == ["P0001"]


def test_existing_event_identity_and_metadata_cannot_be_changed_by_form(setup):
    app, _, events, _ = setup
    title = events.events["EVT-001"].title
    client = app.test_client()
    token = login(client)
    location = start(client, token).headers["Location"]
    page = submit(client, token, app, location, **{"event[eid]": "OTHER", "event[title]": "Override"})
    assert b"Participants imported" in page.data
    assert list(events.events) == ["EVT-001"] and events.events["EVT-001"].title == title


def test_exported_matching_pid_is_ignored_and_identity_is_rechecked(setup):
    app, people, _, _ = setup
    returning(people, name="Someone ELSE")
    item = record()
    item["fields"].update(gender="Female", event_reference="EVT-001")
    item["matches"] = [{"pid": "P0814", "reasons": ["Untrusted file claim"]}]
    client = app.test_client()
    token = login(client)
    location = start(client, token, legacy_excel([item])).headers["Location"]
    assert b"New participant" in client.get(location).data
    assert b"Participants imported" in submit(client, token, app, location).data
    assert people.participants["P0814"].name == "Someone ELSE"
    assert len(people.participants) == 2


def test_duplicate_profiles_require_selection_and_selected_pid_is_verified(setup):
    app, people, events, _ = setup
    first = returning(people)
    people.participants["P0815"] = first.model_copy(update={"pid": "P0815"})
    client = app.test_client()
    token = login(client)
    location = start(client, token).headers["Location"]
    assert b"Participant match needs review" in client.get(location).data
    assert b"Participants imported" not in submit(client, token, app, location).data
    page = client.post(location, data={"csrf_token": token, "include[0]": "1", "match[0]": "P0815"}, follow_redirects=True)
    assert b"Returning participant: P0815" in page.data
    assert b"Participants imported" in submit(client, token, app, location, accepted=[]).data
    assert events.events["EVT-001"].participants == ["POLD", "P0815"]
    assert people.counter == 1


def test_match_changes_must_be_reviewed_before_commit(setup):
    app, people, _, snapshots = setup
    client = app.test_client()
    token = login(client)
    location = start(client, token).headers["Location"]
    returning(people)  # Another upload added this person while the preview was open.
    page = submit(client, token, app, location)
    assert b"matches changed" in page.data and not snapshots.rows
    assert b"Participants imported" in submit(client, token, app, location).data


def test_shared_returning_match_rejects_pid_with_wrong_original_identity(setup):
    _, people, _, _ = setup
    returning(people)
    with pytest.raises(ReviewMatchError):
        find_returning_participant({"_review": {"pid": "P0814", "identity": {"name": "Other PERSON", "dob": "1980-03-25", "representing_country": "AL_TEST"}}}, people)


def test_authentication_csrf_session_ownership_and_feature_flag(setup):
    app, _, _, _ = setup
    first, second = app.test_client(), app.test_client()
    assert first.get("/imports/word/import").status_code == 302
    token = login(first)
    assert start(first, "bad-csrf").status_code == 400
    location = start(first, token).headers["Location"]
    other_token = login(second)
    assert second.get(location).status_code == 404
    assert second.post(location, data={"csrf_token": other_token, "import_now": "1", "include[0]": "1"}).status_code == 404
    assert first.get(location).headers["Cache-Control"] == "no-store"
    app.config["WORD_EXTRACTION_ENABLED"] = False
    assert first.get("/imports/word/import").status_code == 404
    assert first.get(location).status_code == 404


def test_simultaneous_resubmissions_commit_the_batch_only_once(setup):
    app, people, events, snapshots = setup
    owner = app.test_client()
    token = login(owner)
    location = start(owner, token).headers["Location"]
    cookie = owner.get_cookie("session").value
    clients = [app.test_client(), app.test_client()]
    for client in clients:
        client.set_cookie("session", cookie)
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda client: client.post(location, data={"csrf_token": token, "include[0]": "1", "import_now": "1"}), clients))
    assert all(response.status_code in (200, 302) for response in responses)
    assert people.counter == 2 and len(people.participants) == 1
    assert snapshots.writes == 1
    assert events.events["EVT-001"].participants == ["POLD", "P0001"]


def test_database_failure_blocks_import_instead_of_treating_everyone_as_new(setup, monkeypatch):
    app, people, _, snapshots = setup
    client = app.test_client()
    token = login(client)
    def offline():
        raise ServerSelectionTimeoutError("offline")
    monkeypatch.setattr(word_import, "load_match_context", offline)
    assert start(client, token).status_code == 503
    assert not people.participants and not snapshots.rows


@pytest.mark.parametrize("data,filename", [(b"bad", "bad.xlsx"), (b"Name,Country CID\n", "empty.csv"), (b"Name,Name,Country CID\nAna,Ana,AL_TEST", "duplicate.csv"), (b"Name,Country CID\nAna,AL_TEST,extra", "extra.csv"), (b"Name,Country CID\nAna,\xff", "encoding.csv"), (b"old", "file.xls")])
def test_invalid_tables_are_rejected(data, filename):
    with pytest.raises(WordImportError):
        read_export(data, filename)


def test_excel_formula_cells_are_rejected_without_execution():
    book = Workbook()
    book.active.append(["Name", "Country CID"])
    book.active.append(["=1+1", "AL_TEST"])
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(WordImportError, match="formulas"):
        read_export(stream.getvalue(), "formula.xlsx")


def test_atomic_roster_append_and_partial_snapshot_repository_updates():
    calls = []
    class Collection:
        def find_one_and_update(self, query, update, **kwargs):
            calls.append((query, update, kwargs))
            return {"eid": "E", "title": "Keep", "participants": ["POLD", "PNEW"]}
        def update_one(self, query, update, **kwargs):
            calls.append((query, update, kwargs))
    event_repo = EventRepository.__new__(EventRepository)
    snapshot_repo = ParticipantEventRepository.__new__(ParticipantEventRepository)
    event_repo.collection = snapshot_repo.collection = Collection()
    event_repo.add_participants("E", ["PNEW"])
    assert calls[0][1] == {"$addToSet": {"participants": {"$each": ["PNEW"]}}}
    snapshot_repo.upsert_partial("PNEW", "E", {"travel_doc_number": "001"}, [{"file": "source.csv"}])
    assert calls[1][1]["$set"] == {"participant_id": "PNEW", "event_id": "E", "travel_doc_number": "001"}
    assert "bank_name" not in calls[1][1]["$set"]
