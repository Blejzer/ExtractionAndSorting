from docx import Document

from utils.docx_parser import parse_docx


def test_parse_docx_pattern_1(tmp_path):
    file_path = tmp_path / "form.docx"
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    table.cell(0, 0).text = "NAME & LAST NAME"
    table.cell(0, 1).text = "Dragan Mektic"
    table.cell(0, 2).text = "DOB"
    table.cell(0, 3).text = "Dec/24/1956"
    table.cell(1, 0).text = "GENDER"
    table.cell(1, 1).text = "Male"
    table.cell(1, 2).text = "COUNTRY"
    table.cell(1, 3).text = "BiH"
    doc.save(file_path)

    participants = parse_docx(str(file_path))

    assert len(participants) == 1
    assert participants[0]["name"] == "Dragan Mektic"
    assert participants[0]["dob"] == "1956-12-24"
    assert participants[0]["gender"] == "Male"


def test_parse_docx_pattern_3(tmp_path):
    file_path = tmp_path / "list.docx"
    doc = Document()
    doc.add_paragraph("4. Makedonija")
    doc.add_paragraph("1. Dragi ZLATANOVSKI (13. 2. 1965.)")
    doc.save(file_path)

    participants = parse_docx(str(file_path))

    assert len(participants) == 1
    assert participants[0]["name"] == "Dragi ZLATANOVSKI"
    assert participants[0]["dob"] == "1965-02-13"
    assert participants[0]["representing_country"] == "Makedonija"
