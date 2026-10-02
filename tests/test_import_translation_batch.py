from collections import Counter
from threading import Lock
from time import sleep

import pandas as pd

from services.imports import lookup_builders as builders


def test_translation_deduplicates_values_and_bounds_concurrency(monkeypatch):
    lock = Lock()
    calls = Counter()
    active = 0
    peak = 0

    def fake_translate(text, language):
        nonlocal active, peak
        assert language == "en"
        with lock:
            calls[text] += 1
            active += 1
            peak = max(peak, active)
        sleep(0.01)
        with lock:
            active -= 1
        return f"English: {text}"

    monkeypatch.setattr(builders, "translate", fake_translate)
    records = [{"organization": f"org {i}", "unit": "shared", "rank": "", "travel_doc_type": "Passport"} for i in range(12)]
    builders.translate_participant_records(records)
    assert len(calls) == 13
    assert all(count == 1 for count in calls.values())
    assert 1 < peak <= builders.TRANSLATION_WORKERS
    assert all(record["unit"] == "English: shared" for record in records)
    assert all(record["rank"] == "" and record["travel_doc_type"] == "Passport" for record in records)
    builders.translate_participant_records([{"organization": "shared"}])
    assert calls["shared"] == 2  # No participant text is cached between imports.


def test_deferred_lookup_and_aliased_entries(monkeypatch):
    calls = []
    monkeypatch.setattr(builders, "translate", lambda text, lang: calls.append(text) or f"EN:{text}")
    frame = pd.DataFrame([{"Name": "Jane", "Middle name": "Ann", "Last name": "Doe", "Organization": "Org"}])
    lookup = builders.build_lookup_main_online(frame, translate_fields=False)
    assert len(lookup) == 2
    assert calls == []
    assert all(record["organization"] == "Org" for record in lookup.values())
    builders.translate_participant_records(list(lookup.values()))
    assert calls == ["Org"]
    assert all(record["organization"] == "EN:Org" for record in lookup.values())


def test_translation_retains_existing_offline_fallback(monkeypatch):
    from utils import translation

    def unavailable(*args, **kwargs):
        raise ConnectionError("translation unavailable")

    monkeypatch.setattr(translation.requests, "get", unavailable)
    records = [{"organization": "unknown original text", "bio_short": "bonjour tout le monde"}]
    builders.translate_participant_records(records)
    assert records[0]["organization"] == "unknown original text"
    assert records[0]["bio_short"] == "Hello everyone"


def test_preview_translates_all_participants_and_excludes_staff(tmp_path, monkeypatch):
    from io import BytesIO
    from openpyxl import load_workbook
    from openpyxl.worksheet.table import Table
    from utils.participants import initialize_cache
    initialize_cache(None)
    from tests.test_import_service_gender import _workbook_bytes_with_gender
    import services.import_service_v2 as importer

    workbook = load_workbook(BytesIO(_workbook_bytes_with_gender("Female")))
    online = workbook["MAIN ONLINE"]
    headers = {cell.value: cell.column for cell in online[1]}
    for index, table_name in enumerate(("tableInst", "tblFac", "tblTech"), start=1):
        staff = workbook.create_sheet(f"Staff {index}")
        staff.append(["Name and Last Name", "Grade"])
        staff.append([f"Staff Person{index}", 1])
        staff.add_table(Table(displayName=table_name, ref="A1:B2"))
        values = [cell.value for cell in online[2]]
        values[headers["Name"] - 1] = "Staff"
        values[headers["Last name"] - 1] = f"Person{index}"
        values[headers["Organization"] - 1] = f"staff-only {index}"
        online.append(values)
    from openpyxl.utils import get_column_letter
    online.tables["ParticipantsList"].ref = f"A1:{get_column_letter(online.max_column)}5"
    path = tmp_path / "participants-and-staff.xlsx"
    workbook.save(path)
    calls = []
    monkeypatch.setattr(builders, "translate", lambda text, lang: calls.append(text) or f"EN:{text}")
    monkeypatch.setattr(importer, "PREVIEW_PARTICIPANT_LOOKUP", False)
    result = importer.parse_for_commit(str(path))
    assert len(result["attendees"]) == 1
    assert result["attendees"][0]["name"] == "John DOE"
    assert result["attendees"][0]["organization"] == "EN:Ministry"
    assert not any(text.startswith("staff-only") for text in calls)
