"""Capitalization matching across Master Tracker parsing, review and upload.

All participant values are fictional; no live database is required.
"""

from contextlib import nullcontext
from datetime import datetime

from openpyxl import Workbook
from openpyxl.worksheet.table import Table
import pytest

from domain.models.participant import Participant
from repositories.participant_repository import ParticipantRepository
from services.imports.participant_review import (
    ReviewMatchError, annotate_participant_reviews, find_returning_participant,
)
from services.upload_service import upload_preview_data
from tests.test_participant_repository import DummyCollection
from tests.test_upload_service import (
    FakeEventRepo, FakeParticipantEventRepo, FakeParticipantRepo,
    _base_event, _base_participant,
)
from utils.names import participant_name_key
from utils.participants import ParticipantLookupCache, initialize_cache


@pytest.fixture(autouse=True)
def isolated_upload_session(monkeypatch):
    import services.upload_service as uploads

    class Session:
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def start_transaction(self):
            return nullcontext()

    monkeypatch.setattr(uploads.mongodb, "start_session", Session, raising=False)
    monkeypatch.setattr(uploads, "refresh_participant_cache", lambda: None)


class LegacyRepo(FakeParticipantRepo):
    """Use the real repository matcher with a small in-memory persistence layer."""

    def __init__(self, *, dob=None, name="JANE SAMPLE", country="HR"):
        super().__init__()
        self.country_reads = []
        person = Participant.model_validate(
            _base_participant(pid="P1234", name=name, representing_country=country,
                              gender="Female", dob=dob, pob=None),
            context={"allow_missing_dob": True},
        )
        self.participants[person.pid] = person

    def find_by_country(self, cid):
        self.country_reads.append(cid)
        return [person for person in self.participants.values()
                if person.representing_country == cid]

    def find_by_pid(self, pid):
        return self.participants.get(pid)

    def find_by_name_dob_and_representing_country_cid(self, **kwargs):
        repository = ParticipantRepository.__new__(ParticipantRepository)
        repository.collection = DummyCollection([
            person.to_mongo() for person in self.participants.values()
        ])
        return repository.find_by_name_dob_and_representing_country_cid(**kwargs)


@pytest.mark.parametrize("name", ["Jane SAMPLE", "JANE SAMPLE", " jane   sample "])
@pytest.mark.parametrize("stored_dob", [None, datetime(1980, 2, 10)])
def test_cache_matches_capitalization_and_missing_legacy_dob(name, stored_dob):
    repo = LegacyRepo(dob=stored_dob)
    cache = ParticipantLookupCache(repo)
    for _ in range(2):
        person = cache.find_by_display_name_country_and_dob(
            name_display=name, country_name="Croatia",
            dob_source="1980-02-10", representing_country="HR",
        )
        assert person is not None and person.pid == "P1234"
    assert repo.country_reads == ["HR"]
    assert repo.participants["P1234"].name == "JANE SAMPLE"


def test_cache_rejects_country_and_known_dob_conflicts():
    cache = ParticipantLookupCache(LegacyRepo(dob=datetime(1980, 2, 10)))
    assert cache.find_by_display_name_country_and_dob(
        name_display="Jane SAMPLE", country_name="Croatia",
        dob_source="1981-02-10", representing_country="HR",
    ) is None
    assert cache.find_by_display_name_country_and_dob(
        name_display="Jane SAMPLE", country_name="Bosnia and Herzegovina",
        dob_source="1980-02-10", representing_country="BA",
    ) is None


def test_name_key_preserves_accents_punctuation_and_order():
    assert participant_name_key("  JANE   SAMPLE ") == participant_name_key("Jane SAMPLE")
    assert participant_name_key("Žana SAMPLE") != participant_name_key("Zana SAMPLE")
    assert participant_name_key("Jane (A.) SAMPLE") != participant_name_key("Jane A. SAMPLE")
    assert participant_name_key("Jane SAMPLE") != participant_name_key("Sample JANE")


def _master_tracker(path):
    book = Workbook()
    main = book.active
    main.title = "Participants"
    main["A1"] = "PFE26M1 Fictional training"
    main["A2"] = "JUNE 1 - 3 - Zagreb"
    costs = book.create_sheet("COST Overview")
    costs.append(["GRAND TOTAL", 1000])
    names = book.create_sheet("List")
    names.append(["Name (Latin)", "Position", "Phone", "Email"])
    names.append(["SAMPLE, Jane", "Inspector", "", ""])
    names.add_table(Table(displayName="ParticipantsLista", ref="A1:D2"))
    online = book.create_sheet("MAIN ONLINE")
    online.append(["Name", "Last name", "Gender", "Date of Birth (DOB)",
                   "Place Of Birth (POB)", "Country of Birth", "Citizenship(s)"])
    online.append(["Jane", "Sample", "Female", datetime(1980, 2, 10),
                   "Example City", "Croatia", "Croatia"])
    online.add_table(Table(displayName="ParticipantsList", ref="A1:G2"))
    country = book.create_sheet("Cro")
    country.append(["Name and Last Name", "Grade (0 - BL, 1 - Pass, 2 - Excel)"])
    country.append(["Jane SAMPLE", 1])
    country.add_table(Table(displayName="tableCro", ref="A1:B2"))
    book.save(path)


def test_master_tracker_attaches_existing_pid_then_review_fills_birth_details(tmp_path, monkeypatch):
    import services.import_service_v2 as imports

    repo = LegacyRepo()
    initialize_cache(repo)
    monkeypatch.setattr(imports, "_participant_repo", repo)
    monkeypatch.setattr(imports, "get_country_cid_by_name", lambda value: "HR" if value else None)
    monkeypatch.setattr(imports, "resolve_country_flexible",
                        lambda value: {"cid": "HR", "country": "Croatia"} if value else None)
    monkeypatch.setattr(imports, "translate_participant_records", lambda records: None)
    path = tmp_path / "PFE26M1.xlsx"
    _master_tracker(path)
    result = imports.parse_for_commit(str(path), preview_only=False)
    record = result["preview"]["participants"][0]
    assert record["pid"] == "P1234"
    assert record["name"] == "Jane SAMPLE"
    assert record["dob"] == "1980-02-10"
    assert record["pob"] == "Example City"

    reviewed = annotate_participant_reviews([record], repo)[0]
    assert reviewed["_review"]["pid"] == "P1234"
    assert {"dob", "pob"} <= set(reviewed["_review"]["accepted_fields"])
    assert reviewed["_changes"]["name"] == {"stored": "JANE SAMPLE", "file": "Jane SAMPLE"}
    # The saved review is checked again by PID with the original, mixed-case name.
    assert find_returning_participant(reviewed, repo).pid == "P1234"


@pytest.mark.parametrize("use_file_name", [False, True])
def test_review_and_upload_keep_pid_and_history_with_case_variant(use_file_name):
    repo = LegacyRepo()
    incoming = _base_participant(name="Jane SAMPLE", gender="Female",
                                dob="1980-02-10", pob="Example City")
    record = annotate_participant_reviews([incoming], repo)[0]
    assert record["_review"]["pid"] == "P1234"
    record["_review"]["accepted_fields"] = ["dob", "pob"] + (["name"] if use_file_name else [])
    snapshots, events = FakeParticipantEventRepo(), FakeEventRepo()
    previous = object()
    snapshots.snapshots.append(previous)
    result = upload_preview_data(
        {"event": _base_event(), "participants": [record]},
        participant_repo=repo, event_repo=events, participant_event_repo=snapshots,
    )
    person = repo.participants["P1234"]
    assert list(repo.participants) == ["P1234"] and repo.counter == 1
    assert person.name == ("Jane SAMPLE" if use_file_name else "JANE SAMPLE")
    assert person.dob == datetime(1980, 2, 10) and person.pob == "Example City"
    assert result["event"].participants == ["P1234"]
    assert snapshots.snapshots[0] is previous
    assert snapshots.snapshots[1].participant_id == "P1234"


def test_saved_review_still_rejects_conflicting_identity():
    repo = LegacyRepo()
    record = annotate_participant_reviews([
        _base_participant(name="Jane SAMPLE", dob="1980-02-10"),
    ], repo)[0]
    repo.participants["P1234"] = repo.participants["P1234"].model_copy(
        update={"dob": datetime(1981, 2, 10)},
    )
    with pytest.raises(ReviewMatchError, match="match changed"):
        find_returning_participant(record, repo)
