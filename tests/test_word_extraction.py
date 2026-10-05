"""Synthetic documents only: no uploaded participant information in fixtures."""

import csv
from datetime import datetime
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile

from flask import Flask
from openpyxl import load_workbook
from pymongo.errors import ServerSelectionTimeoutError
import pytest

from routes import word_extraction as routes
from services import word_draft_store as drafts
from services.word_export_service import EXPORT_COLUMNS, WordExportError, export_csv, export_xlsx
from services.word_extraction_service import FIELDS, WordExtractionError, extract_docx, extract_files, parse_date
from services.word_matching_service import WordMatchContext, load_match_context


def paragraph(text):
    runs = "<w:tab/>".join(f"<w:t>{escape(part)}</w:t>" for part in text.split("\t"))
    return f"<w:p><w:r>{runs}</w:r></w:p>"


def table(rows):
    return "<w:tbl>" + "".join("<w:tr>" + "".join(
        "<w:tc>" + "".join(paragraph(line) for line in cell.split("\n")) + "</w:tc>"
        for cell in row) + "</w:tr>" for row in rows) + "</w:tbl>"


def docx(*blocks):
    xml = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + "".join(blocks) + '</w:body></w:document>'
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", xml)
    return stream.getvalue()


def form(name="Ana TEST", country="Albania", dob="25.03.1980."):
    return docx(table([
        ["COUNTRY:", country], ["NAME & LAST NAME:", name, "DOB:", dob, "POB:", "Tirana"],
        ["PASSPORT NUMBER:", "001234", "GENDER:", "F"],
        ["INSTITUTION:", "Test institution", "POSITION:", "Specialist"],
        ["SHORT PROFESSIONAL BIOGRAPHY:", "First paragraph.\nSecond paragraph."],
    ]))


COUNTRIES = [
    {"cid": "AL_TEST", "country": "Albania, Europe & Eurasia", "iso": "AL"},
    {"cid": "BA_TEST", "country": "Bosnia and Herzegovina, Europe & Eurasia", "iso": "BA"},
    {"cid": "MK_TEST", "country": "North Macedonia, Europe & Eurasia", "iso": "MK"},
    {"cid": "RS_TEST", "country": "Serbia, Europe & Eurasia", "iso": "RS"},
    {"cid": "XK_TEST", "country": "Kosovo, Europe & Eurasia", "iso": "XK"},
]
PARTICIPANTS = [{"pid": "PTEST1", "name": "Ana TEST", "representing_country": "AL_TEST",
                 "dob": datetime(1980, 3, 25), "email": "ana@example.test"}]


def record(name="Ana TEST", country="Albania", dob="1980-03-25"):
    return {"fields": {field: "" for field in FIELDS} | {"name": name, "country_label": country, "dob": dob},
            "source_file": "test.docx", "sources": ["test.docx: Table 1"], "evidence": {},
            "warnings": [], "matches": [], "match_status": "Not checked"}


def test_registration_table_preserves_fields_and_multiline_bio():
    result = extract_docx(form(), "PFE14M1.docx")
    assert len(result) == 1
    fields = result[0]["fields"]
    assert fields["name"] == "Ana TEST"
    assert fields["dob"] == "1980-03-25"
    assert fields["country_label"] == "Albania"
    assert fields["representing_country"] == ""
    assert fields["gender"] == "Female"
    assert fields["bio_short"] == "First paragraph.\nSecond paragraph."
    assert fields["event_reference"] == "PFE14M1"
    assert fields["travel_doc_number"] == "001234"
    assert fields["travel_doc_type"] == "Passport"


def test_multiple_forms_and_mixed_date_formats_require_review():
    data = docx(table([["Name:", "Ana TEST", "DOB:", "07/30/1980"]]),
                table([["Name:", "Bora TEST", "DOB:", "30/07/1981"]]),
                table([["Name:", "Cora TEST", "DOB:", "2/10/1982"]]))
    result = extract_docx(data, "forms.docx")
    assert [r["fields"]["dob"] for r in result] == ["1980-07-30", "1981-07-30", ""]
    assert result[2]["fields"]["dob_raw"] == "2/10/1982"
    assert any("ambiguous" in warning for warning in result[2]["warnings"])
    assert extract_docx(data, "forms.docx", "dmy")[2]["fields"]["dob"] == "1982-10-02"


def test_repeated_name_labels_in_one_table_start_separate_entries():
    data = docx(table([["Name:", "Ana TEST", "DOB:", "25.03.1980"],
                       ["Name:", "Bora TEST", "DOB:", "26.03.1981"]]))
    result = extract_docx(data, "forms.docx")
    assert [(r["fields"]["name"], r["fields"]["dob"]) for r in result] == [
        ("Ana TEST", "1980-03-25"), ("Bora TEST", "1981-03-26")]


def test_tabbed_lists_keep_column_boundaries_and_date_evidence():
    data = docx(paragraph("Albania"), paragraph("Ana TEST\t07/30/1980\tMajor\tDirector"),
                paragraph("Bora TEST\t2/10/1982\tOfficer\tSpecialist"))
    result = extract_docx(data, "list.docx")
    assert len(result) == 2
    assert result[1]["fields"]["dob"] == "1982-02-10"
    assert result[0]["fields"]["rank"] == "Major"
    assert result[0]["fields"]["position"] == "Director"


def test_numbered_biographies_and_country_sections():
    data = docx(paragraph("Montenegro"), paragraph("1. Ana TEST, rođena 25.03.1980. u Podgorici."),
                paragraph("An additional biography paragraph."), paragraph("Makedonija:"),
                paragraph("Let Skopje-Zagreb, 21.4."), paragraph("1. Bora TEST (26.03.1981.)"),
                paragraph("Srbija:"), paragraph("Automobilima iz Beograda"), paragraph("1. Cora TEST"))
    result = extract_docx(data, "PFE13M3.docx")
    assert len(result) == 3
    assert result[0]["fields"]["bio_short"].endswith("An additional biography paragraph.")
    assert result[1]["fields"]["country_label"] == "Makedonija"
    assert result[1]["fields"]["travel_notes"].startswith("Let Skopje")
    assert result[2]["fields"]["country_label"] == "Srbija"
    assert result[2]["fields"]["dob"] == ""
    assert result[2]["fields"]["travel_notes"].startswith("Automobilima")


def test_prose_and_supplemental_roster_combine_only_known_identity():
    data = docx(paragraph("Capt. ANA TEST - PK#1234"), paragraph("DOB: 25.03.1980"),
                paragraph("Position: Specialist"), paragraph("Passport No: 001234"),
                table([["CASE ID", "NAME", "DOB", "UNIT NAME"], ["V001", "TEST, Ana", "03/25/1980", "Investigations"]]))
    result = extract_docx(data, "Kosovo.docx")
    assert len(result) == 1
    fields = result[0]["fields"]
    assert fields["name"] == "ANA TEST"
    assert fields["unit"] == "Investigations"
    assert fields["case_id"] == "V001"
    assert fields["service_number"] == "1234"
    assert fields["travel_doc_number"] == "001234"
    assert fields["country_label"] == ""  # No inference from filename or institution.
    assert len(result[0]["evidence"]["name"]) == 2


def test_same_name_and_dob_with_different_countries_stay_separate():
    data = docx(table([["Country:", "Albania"], ["Name:", "Ana TEST", "DOB:", "25.03.1980"]]),
                table([["Country:", "Serbia"], ["Name:", "Ana TEST", "DOB:", "25.03.1980"]]))
    assert len(extract_docx(data, "forms.docx")) == 2


def test_conflicting_source_values_preserved_for_review():
    data = docx(paragraph("Name: Ana TEST"), paragraph("DOB: 25.03.1980"),
                paragraph("Position: Specialist"), paragraph("Position: Director"))
    result = extract_docx(data, "form.docx")[0]
    assert result["fields"]["position"] == "Specialist"
    assert [item["value"] for item in result["evidence"]["position"]] == ["Specialist", "Director"]
    assert any("Different position" in item for item in result["warnings"])


def test_country_before_name_in_paragraph_form():
    data = docx(paragraph("Country: Republic of Serbia"), paragraph("Name:"), paragraph("Ana TEST"), paragraph("DOB: 25.03.1980"))
    result = extract_docx(data, "form.docx")[0]
    assert result["fields"]["country_label"] == "Republic of Serbia"
    assert result["fields"]["name"] == "Ana TEST"


def test_partial_batch_errors_and_cross_file_duplicates():
    result = extract_files([("one.docx", form()), ("two.docx", form()), ("old.doc", b"old"), ("bad.docx", b"bad")])
    assert len(result["records"]) == 2
    assert len(result["file_errors"]) == 2
    assert all(any("another uploaded file" in warning for warning in item["warnings"]) for item in result["records"])


def test_entry_limits_stop_parsing_before_allocating_a_large_batch(monkeypatch):
    monkeypatch.setattr("services.word_extraction_service.MAX_RECORDS", 2)
    with pytest.raises(WordExtractionError, match="document contains more than 2"):
        extract_docx(docx(paragraph("Name: Ana TEST"), paragraph("Name: Bora TEST"), paragraph("Name: Cora TEST")), "large.docx")
    with pytest.raises(WordExtractionError, match="batch contains more than 2"):
        extract_files([("one.docx", form()), ("two.docx", form()), ("three.docx", form())])


@pytest.mark.parametrize("value,expected", [
    ("2/10/1982", ""), ("30/07/1980", "1980-07-30"), ("07/30/1980", "1980-07-30"),
    ("25. 03. 1980.", "1980-03-25"), ("March 25, 1980", "1980-03-25"),
    ("1980-03-25T00:00:00Z", "1980-03-25"), (datetime(1980, 3, 25), "1980-03-25"),
    ("NOV/261/1971", ""), ("31.02.1980", ""), ("03/25/80", ""), ("March", ""), ("March 1980", ""),
])
def test_dates_never_guess_ambiguous_or_invalid_values(value, expected):
    assert parse_date(value) == expected


@pytest.mark.parametrize("content", [b"not a zip", docx(paragraph("No names here")),
                                     docx(paragraph("&ignored;"))])
def test_unsupported_documents_return_extraction_errors(content):
    with pytest.raises(WordExtractionError):
        extract_docx(content, "invalid.docx")


def test_xml_entities_and_oversized_documents_rejected(monkeypatch):
    stream = BytesIO()
    with ZipFile(stream, "w") as archive:
        archive.writestr("word/document.xml", '<!DOCTYPE x [<!ENTITY a "expanded">]><x>&a;</x>')
    with pytest.raises(WordExtractionError, match="entity"):
        extract_docx(stream.getvalue(), "entities.docx")
    monkeypatch.setattr("services.word_extraction_service.MAX_XML_BYTES", 10)
    with pytest.raises(WordExtractionError, match="limit"):
        extract_docx(form(), "large.docx")


def test_authoritative_country_cid_and_identity_match():
    context = WordMatchContext(PARTICIPANTS, COUNTRIES)
    item = record(name="TÉST, Ana")
    context.check(item)
    assert item["fields"]["representing_country"] == "AL_TEST"
    assert item["match_status"] == "Exact match"
    assert item["matches"][0]["pid"] == "PTEST1"
    assert context.resolve_country("Bosnia and Herzegovina – Republic of Srpska Entity") == "BA_TEST"
    assert context.resolve_country("Makedonija") == "MK_TEST"


@pytest.mark.parametrize("changes", [{"dob": ""}, {"dob": "1981-03-25"}, {"country_label": ""}, {"country_label": "Serbia"}, {"name": "Anna TEST"}])
def test_missing_conflicting_or_similar_identity_is_only_possible(changes):
    item = record()
    item["fields"].update(changes)
    WordMatchContext(PARTICIPANTS, COUNTRIES).check(item)
    assert item["match_status"] == "Possible match"
    assert not item["matches"][0]["exact"]


def test_duplicates_and_unknown_country_remain_reviewable():
    context = WordMatchContext(PARTICIPANTS + [PARTICIPANTS[0] | {"pid": "PTEST2"}], COUNTRIES)
    item = record()
    context.check(item)
    assert item["match_status"] == "Multiple exact matches"
    item["fields"]["representing_country"] = "INVALID_CID"
    context.check(item)
    assert item["match_status"] == "Review required"
    assert item["matches"] == []
    item = record(name="Unrelated Person", country="")
    context.check(item)
    assert item["match_status"] == "Review required"
    item["fields"]["country_label"] = "Serbia"
    context.check(item)
    assert item["match_status"] == "No candidate found"


def test_match_context_uses_only_read_queries_and_raw_legacy_documents(monkeypatch):
    import config.database as database
    calls = []
    class ReadOnlyCollection:
        def __init__(self, name):
            self.name = name
        def find(self, query, projection):
            calls.append((self.name, query, projection))
            return PARTICIPANTS if self.name == "participants" else COUNTRIES
    class ReadOnlyDatabase:
        def collection(self, name):
            return ReadOnlyCollection(name)
    monkeypatch.setattr(database, "mongodb", ReadOnlyDatabase())
    context = load_match_context()
    item = record()
    context.check(item)
    assert item["match_status"] == "Exact match"
    assert [call[0] for call in calls] == ["participants", "countries"]
    assert all(call[1] == {} and call[2]["_id"] == 0 for call in calls)


def test_exports_share_columns_preserve_zeros_and_neutralize_formulas():
    item = record()
    item["fields"].update(name="=1+1", phone="+1234", travel_doc_number="001234", bio_short="First\nSecond")
    data = export_xlsx([item])
    book = load_workbook(BytesIO(data), data_only=False)
    sheet = book.active
    assert list(next(sheet.values)) == list(EXPORT_COLUMNS.values())
    row = dict(zip(EXPORT_COLUMNS, list(sheet.values)[1]))
    assert row["travel_doc_number"] == "001234"
    assert row["bio_short"] == "First\nSecond"
    assert row["name"] == "=1+1"
    assert all(cell.data_type != "f" for cell in sheet[2])
    csv_rows = list(csv.reader(StringIO(export_csv([item]).decode("utf-8-sig"))))
    assert csv_rows[0] == list(EXPORT_COLUMNS.values())
    csv_row = dict(zip(EXPORT_COLUMNS, csv_rows[1]))
    assert csv_row["name"] == "'=1+1"
    assert csv_row["phone"] == "'+1234"
    assert csv_row["travel_doc_number"] == "001234"


def test_excel_does_not_silently_truncate_long_source_text():
    item = record()
    item["fields"]["bio_short"] = "x" * 32768
    with pytest.raises(WordExportError, match="Download CSV"):
        export_xlsx([item])
    assert item["fields"]["bio_short"] in export_csv([item]).decode("utf-8-sig")


@pytest.fixture
def app(tmp_path, monkeypatch):
    application = Flask(__name__, template_folder=str(Path(__file__).resolve().parents[1] / "templates"))
    application.config.update(TESTING=True, SECRET_KEY="word-test", WORD_EXTRACTION_DIR=str(tmp_path / "drafts"))
    application.add_url_rule("/", endpoint="main.show_home", view_func=lambda: "home")
    for endpoint in ("participants.show_participants", "events.show_events", "imports.upload_form", "auth.login", "auth.logout"):
        application.add_url_rule("/" + endpoint, endpoint=endpoint, view_func=lambda: "stub")
    application.add_url_rule("/participants/<pid>", endpoint="participants.participant_detail", view_func=lambda pid: pid)
    application.register_blueprint(routes.word_extraction_bp)
    monkeypatch.setattr(routes, "load_match_context", lambda: WordMatchContext(PARTICIPANTS, COUNTRIES))
    return application


def login(client, username="tester"):
    with client.session_transaction() as session:
        session["username"] = username
    client.get("/imports/word")
    with client.session_transaction() as session:
        return session["word_extraction_csrf"]


def upload(client, token, content=None, filename="form.docx"):
    return client.post("/imports/word", data={"csrf_token": token, "files": (BytesIO(content or form()), filename)}, content_type="multipart/form-data")


def test_authenticated_upload_edit_export_and_clear(app):
    client = app.test_client()
    assert client.get("/imports/word").status_code == 302
    token = login(client)
    assert upload(client, "wrong-token").status_code == 400
    response = upload(client, token)
    assert response.status_code == 302
    location = response.headers["Location"]
    page = client.get(location)
    assert b"Exact match" in page.data and b"PTEST1" in page.data
    assert page.headers["Cache-Control"] == "no-store"
    exported = client.post(location + "/export.xlsx", data={"csrf_token": token, "r0_position": "Updated specialist"})
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.data)).active
    values = dict(zip(EXPORT_COLUMNS, list(sheet.values)[1]))
    assert values["position"] == "Updated specialist"
    assert values["matching_pids"] == "PTEST1"
    assert len(list(Path(app.config["WORD_EXTRACTION_DIR"]).glob("*.json"))) == 1
    assert client.post(location + "/clear", data={"csrf_token": token}).status_code == 302
    assert client.get(location).status_code == 410
    assert not list(Path(app.config["WORD_EXTRACTION_DIR"]).glob("*.json"))


def test_country_can_be_selected_for_all_unstated_entries(app):
    client = app.test_client()
    token = login(client)
    response = upload(client, token, form(country=""))
    location = response.headers["Location"]
    assert b"Possible match" in client.get(location).data
    assert client.post(location + "/check", data={"csrf_token": token, "missing_country": "AL_TEST"}).status_code == 302
    assert b"Exact match" in client.get(location).data


def test_another_session_cannot_view_edit_export_or_clear_a_draft(app):
    first, second = app.test_client(), app.test_client()
    location = upload(first, login(first)).headers["Location"]
    token = login(second)  # Even the same username must not inherit the other session's upload.
    assert second.get(location).status_code == 404
    for operation in ("check", "export.csv", "clear"):
        assert second.post(location + "/" + operation, data={"csrf_token": token}).status_code == 404
    assert first.get(location).status_code == 200


def test_unavailable_database_still_allows_extraction_and_export(app, monkeypatch):
    def unavailable():
        raise ServerSelectionTimeoutError("offline")
    monkeypatch.setattr(routes, "load_match_context", unavailable)
    client = app.test_client()
    token = login(client)
    location = upload(client, token).headers["Location"]
    page = client.get(location)
    assert b"Not checked" in page.data
    assert b"No candidate found" not in page.data.split(b"<tbody>")[1]
    response = client.post(location + "/export.csv", data={"csrf_token": token})
    assert response.status_code == 200
    rows = list(csv.reader(StringIO(response.data.decode("utf-8-sig"))))
    assert dict(zip(EXPORT_COLUMNS, rows[1]))["match_status"] == "Not checked"


def test_invalid_dates_missing_files_and_disabled_tool(app):
    client = app.test_client()
    token = login(client)
    assert client.post("/imports/word", data={"csrf_token": token}).status_code == 400
    location = upload(client, token).headers["Location"]
    assert client.post(location + "/check", data={"csrf_token": token, "r0_dob": "31.02.1980"}).status_code == 400
    assert b"Exact match" in client.get(location).data  # Failed edits did not overwrite the draft.
    app.config["WORD_EXTRACTION_ENABLED"] = False
    assert client.get("/imports/word").status_code == 404
    assert client.post(location + "/export.xlsx", data={"csrf_token": token}).status_code == 404


def test_expired_drafts_are_removed_and_not_stored_in_session(app):
    client = app.test_client()
    location = upload(client, login(client)).headers["Location"]
    path = next(Path(app.config["WORD_EXTRACTION_DIR"]).glob("*.json"))
    assert path.stat().st_mode & 0o777 == 0o600
    with client.session_transaction() as session:
        assert "Ana TEST" not in json.dumps(dict(session))
    past = path.stat().st_mtime - drafts.TTL_SECONDS - 1
    os.utime(path, (past, past))
    assert client.get(location).status_code == 410
    assert not path.exists()
