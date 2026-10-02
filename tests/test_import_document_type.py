from io import BytesIO

import pytest
from openpyxl import load_workbook


@pytest.mark.parametrize("selection,expected", [("Passport", "Passport"), ("Other", "ID Card")])
@pytest.mark.parametrize("compound_surname", [False, True])
def test_registration_document_survives_position_match(tmp_path, monkeypatch, selection, expected, compound_surname):
    # Initialize the cache explicitly for this test, avoiding the existing
    # module-level cache initialization defect without changing production code.
    from utils.participants import initialize_cache
    initialize_cache(None)
    from tests.test_import_service_gender import _workbook_bytes_with_gender
    import services.import_service_v2 as importer

    workbook = load_workbook(BytesIO(_workbook_bytes_with_gender("Female")))
    first, last = ("Ana", "De Silva") if compound_surname else ("Ana", "Silva")
    display = f"{first} {last}"
    workbook["List"]["A2"] = display
    workbook["Cro"]["A2"] = display
    online = workbook["MAIN ONLINE"]
    headers = {cell.value: cell.column for cell in online[1]}
    online.cell(2, headers["Name"], first)
    online.cell(2, headers["Last name"], last)
    column = online.max_column + 1
    online.cell(1, column, "Traveling document type")
    online.cell(2, column, selection)
    from openpyxl.utils import get_column_letter
    table = online.tables["ParticipantsList"]
    table.ref = f"A1:{get_column_letter(column)}2"
    table.tableColumns = []
    path = tmp_path / "registration.xlsx"
    workbook.save(path)
    monkeypatch.setattr(importer, "PREVIEW_PARTICIPANT_LOOKUP", False)
    result = importer.parse_for_commit(str(path))
    attendee = result["attendees"][0]
    assert attendee["travel_doc_type"] == expected
    assert attendee["position"] == "Analyst"
    assert attendee["dob"]
    assert result["preview"]["participants"][0]["travel_doc_type"] == expected
